"""Python에서 계산한 계좌·ETF·매매 결과를 세 개의 엑셀 시트로 저장한다."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .backtest import BacktestResult


def save_excel(result: BacktestResult, signals: pd.DataFrame, prices: pd.DataFrame, path: Path):
    # 계좌 잔액은 현금과 ETF 평가액으로 나눠 표시한다.
    curve = result.curve
    account = pd.DataFrame({
        "날짜": curve.index, "현금(USD)": curve["cash"].to_numpy(),
        "ETF평가액(USD)": (curve["nav"] - curve["cash"]).to_numpy(),
        "총자산(USD)": curve["nav"].to_numpy(), "일수익률": curve["net_return"].to_numpy(),
        "누적수익률": (curve["nav"] / curve["nav"].iloc[0] - 1).to_numpy(),
        "거래비용(USD)": curve["fee"].to_numpy(), "누적비용(USD)": curve["fee"].cumsum().to_numpy(),
        "매매금액(USD)": curve["traded_notional"].to_numpy(), "회전율": curve["turnover"].to_numpy(),
    })

    # 한 행에 하루의 ETF 하나를 담는다. 신호일과 체결일을 분리한다.
    rows = []
    for signal in signals.itertuples(index=False):
        day, symbol = signal.execution_date, signal.symbol
        price = float(prices.at[day, symbol])
        weight = float(result.holdings.at[day, symbol])
        value = weight * float(curve.at[day, "nav"])
        rows.append({
            "체결일": day, "신호일": signal.signal_date, "ETF": symbol,
            "기대수익률(연율)": signal.expected_annual_return,
            "변동성(연율)": signal.annualized_volatility, "분산(연율)": signal.annualized_variance,
            "목표비중": signal.target_weight, "보유비중": weight,
            "보유수량(조정단위)": value / price, "조정종가(USD)": price,
            "평가액(USD)": value, "베타": signal.beta,
        })
    assets = pd.DataFrame(rows)
    trades = result.trades.rename(columns={
        "signal_date": "신호일", "execution_date": "체결일", "symbol": "ETF", "side": "매매",
        "adjusted_units": "수량(조정단위)", "adjusted_price": "조정종가(USD)",
        "notional": "거래금액(USD)", "fee": "거래비용(USD)", "target_weight": "목표비중",
    })
    trades["매매"] = trades["매매"].replace({"BUY": "매수", "SELL": "매도"})

    sheets = {"일별 계좌": account, "ETF별 계산": assets, "매매 내역": trades}
    notes = {
        "일별 계좌": "총자산은 현금과 ETF 평가액의 합계. 거래비용 차감 후 USD 기준.",
        "ETF별 계산": "252거래일 추정. 신호 다음 거래일 종가에 체결. 수량은 조정가격 기준 분수 단위.",
        "매매 내역": "목표 보유액과 기존 보유액의 차이만 거래. 실제 증권사 주문이 아닌 백테스트 기록.",
    }
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False, startrow=4, freeze_panes=(5, 3 if name == "ETF별 계산" else 1))
            sheet = writer.sheets[name]
            sheet.sheet_view.showGridLines = False
            sheet["A2"], sheet["A3"] = name, notes[name]
            sheet["A2"].font = Font(name="Arial", size=14, bold=True)
            sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")
            sheet.auto_filter.ref = f"A5:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
            sheet.row_dimensions[5].height = 32
            for column in sheet.iter_cols(min_row=5):
                header = column[0]
                sheet.column_dimensions[header.column_letter].width = 22 if len(str(header.value)) > 9 else 17
                header.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                header.fill = PatternFill("solid", fgColor="24476A")
                header.alignment = Alignment(horizontal="center", vertical="center")
                label = str(header.value)
                number_format = "#,##0.00;(#,##0.00);0.00"
                if any(word in label for word in ("수익률", "변동성", "비중", "회전율")):
                    number_format = "0.00%;(0.00%);0.00%"
                elif "분산" in label or "수량" in label:
                    number_format = "0.000000"
                elif "베타" in label or "종가" in label:
                    number_format = "0.0000"
                elif label in ("날짜", "신호일", "체결일"):
                    number_format = "yyyy-mm-dd"
                for cell in column[1:]:
                    cell.font = Font(name="Arial", size=10)
                    cell.number_format = number_format
        source = writer.sheets["일별 계좌"]
        for row, label, url in (
            (2, "Yahoo Finance", "https://finance.yahoo.com/quote/SPY/history/"),
            (3, "Kenneth French", "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/Data_Library/f-f_factors.html"),
        ):
            source.cell(row, 12, "데이터 출처")
            source.cell(row, 13, label).hyperlink = url
            source.column_dimensions["M"].width = 22
